from docx import Document
import json

def parse_CMSTC_sections():
    name = "CMSTC"

    docx = "input/" + name + ".docx"
    txt = "input/" + name + ".txt"

    doc = Document(docx)

    result = []
    with open(txt, "r") as file:
        s = file.readline().strip()
        level_numbers = [None, None, None, None]
        levels = [None, None, None, None]
        while len(s) == 0 or s[0] != "4":
            s = file.readline().strip()
        
        for line in file:
            s = line.strip()
            if len(s) > 1 and s[0].isdigit() and s[1] == ".":
                numbering = s.split()[0]
                curr_levels = numbering.split(".")
                at_level = len(curr_levels) - 2

                # update levels
                levels[at_level] = s[len(numbering) + 1:]
                level_numbers[at_level] = int(curr_levels[-1])

                # add new dict
                if at_level >= 2:
                    tc = {
                        "id": numbering,
                        "feature": levels[0],
                        "subfeature": levels[1],
                        "test": levels[2],
                        "steps": []
                    }
                    if at_level >= 3:
                        tc["subtest"] = levels[3]
                    result.append(tc)

    for i in range(len(doc.tables)):
        table = doc.tables[i]
        curr = result[i]
        keys = [cell.text.strip() for cell in table.rows[0].cells]
        for row in table.rows[1:]:
            values = [cell.text.strip() for cell in row.cells]
            row_dict = dict(zip(keys, values))
            del row_dict["Actual Results"]
            curr["steps"].append(row_dict)
        result[i] = curr

    with open("output/" + name + ".json", "w") as file:
        json.dump(result, file, indent=4)


def parse_PRSTC_sections():
    name = "PRSTC"

    docx = "input/" + name + ".docx"
    txt = "input/" + name + ".txt"

    doc = Document(docx)

    extra_info = {
        "Use case identifier : ",
        "Requirement ID : ",
        "Modules: ",
        "describe : ",
        "Quality characteristics : ",
        "Evaluation criteria : ",
        "Termination conditions : "
        }

    result = []
    with open(txt, "r") as file:
        s = file.readline().strip()
        level_numbers = [None, None, None, None]
        levels = [None, None, None, None]
        while len(s) == 0 or s[0] != "4":
            s = file.readline().strip()
        
        extra = {}
        
        for line in file:
            s = line.strip()
            if len(s) > 1 and s[0].isdigit() and "." in s:
                if 4 <= len(line) - len(s) <= 5:
                    at_level = 0
                elif 7 <= len(line) - len(s) <= 8:
                    at_level = 1
                else:
                    continue

                num = s.split(".")[0]
                levels[at_level] = s[len(num) + 2:]
                level_numbers[at_level] = num

                if at_level >= 1:
                    tc = {
                        "id": ".".join(level_numbers[:at_level + 1]),
                        "feature": levels[0],
                        "test": levels[1],
                        "extra_info": extra,
                        "steps": []
                    }
                    result.append(tc)
            else:
                for info in extra_info:
                    if info in s:
                        key = info[:-3] if info[-3] == " " else info[:-2]
                        extra[key] = s[len(info):]
                        break

    
    for i in range(len(doc.tables)):
        table = doc.tables[i]
        curr = result[i]
        keys = [cell.text.strip() for cell in table.rows[0].cells]
        for row in table.rows[1:]:
            values = [cell.text.strip() for cell in row.cells]
            row_dict = dict(zip(keys, values))
            curr["steps"].append(row_dict)
        result[i] = curr
    
    with open("output/" + name + ".json", "w") as file:
        json.dump(result, file, indent=4)


def parse_SMSTC_sections():
    name = "SMSTC"

    docx = "input/" + name + ".docx"
    txt = "input/" + name + ".txt"

    doc = Document(docx)

    result = []
    with open(txt, "r") as file:
        s = file.readline().strip()
        level_numbers = [None, None, None, None]
        levels = [None, None, None, None]
        while len(s) == 0 or s[0] != "4":
            s = file.readline().strip()
        
        for line in file:
            s = line.strip()
            if len(s) > 1 and s[0].isdigit() and "." in s:
                if 4 <= len(line) - len(s) <= 5:
                    at_level = 0
                elif 7 <= len(line) - len(s) <= 8:
                    at_level = 1
                else:
                    continue

                num = s.split(".")[0]
                levels[at_level] = s[len(num) + 2:]
                level_numbers[at_level] = num

                if at_level >= 1:
                    tc = {
                        "id": ".".join(level_numbers[:at_level + 1]),
                        "feature": levels[0],
                        "test": levels[1],
                        "steps": []
                    }
                    result.append(tc)

    for i in range(len(doc.tables)):
        table = doc.tables[i]
        curr = result[i]
        keys = [cell.text.strip() for cell in table.rows[0].cells]
        for row in table.rows[1:]:
            values = [cell.text.strip() for cell in row.cells]
            row_dict = dict(zip(keys, values))
            if values[1] == values[2]:
                del row_dict["Expected Results"]
            curr["steps"].append(row_dict)
        result[i] = curr

    with open("output/" + name + ".json", "w") as file:
        json.dump(result, file, indent=4)


# parse_CMSTC_sections()
# parse_PRSTC_sections()
parse_SMSTC_sections()



